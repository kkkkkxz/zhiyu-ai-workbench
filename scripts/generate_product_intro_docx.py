# -*- coding: utf-8 -*-
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TITLE = "知语AI智能知识服务平台"
SUBTITLE = "产品介绍"
COMPANY = "展望数科"
PUBLIC_DEMO_URL = os.environ.get("ZWSK_PUBLIC_DEMO_URL", "").strip()
FONT_CN = "微软雅黑"
FONT_EN = "Arial"

BLUE = RGBColor(18, 58, 95)
CYAN = RGBColor(30, 109, 143)
GOLD = RGBColor(176, 132, 46)
TEXT = RGBColor(35, 43, 51)
MUTED = RGBColor(96, 108, 125)


def desktop_dir() -> Path:
    path = Path.home() / "Desktop"
    if path.exists():
        return path
    return Path("C:/Users/admin/Desktop")


def set_run_font(run, size: float = 10.5, bold: bool = False, color: RGBColor = TEXT):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)


def set_style_font(style, size: float):
    style.font.name = FONT_EN
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D7DEE8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        borders.append(border)
    tc_pr.append(borders)


def set_cell_margins(cell, top=110, start=150, bottom=110, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def paragraph(doc, text: str = "", size: float = 10.5, bold: bool = False,
              color: RGBColor = TEXT, align=WD_ALIGN_PARAGRAPH.LEFT,
              before: float = 0, after: float = 5, first_indent: bool = False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.32
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def heading(doc, text: str, level: int = 1):
    size = {1: 17, 2: 13, 3: 11.5}.get(level, 11.5)
    color = {1: BLUE, 2: CYAN, 3: TEXT}.get(level, TEXT)
    p = paragraph(doc, text, size=size, bold=True, color=color, before=10 if level == 1 else 6, after=6)
    if level == 1:
        p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run("• " + text)
    set_run_font(run, size=10.2, color=TEXT)


def add_meta_table(doc):
    rows = [
        ("文档名称", f"{TITLE}{SUBTITLE}"),
        ("适用对象", "政府及企事业单位、企业知识管理团队、客服与运营部门、销售与项目管理团队"),
        ("部署形态", "本地化部署、专有云部署、Docker Compose 或 Windows 原生服务部署"),
        ("当前访问", "Cloudflare Tunnel 公网临时访问，前端代理 /api 到本地后端服务"),
        ("生成日期", datetime.now().strftime("%Y年%m月%d日")),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (key, value) in enumerate(rows):
        left, right = table.cell(i, 0), table.cell(i, 1)
        left.width = Cm(3.2)
        right.width = Cm(12.6)
        left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade_cell(left, "EAF3F8")
        set_cell_borders(left)
        set_cell_borders(right)
        set_cell_margins(left)
        set_cell_margins(right)
        r1 = left.paragraphs[0].add_run(key)
        set_run_font(r1, 10.2, True, BLUE)
        r2 = right.paragraphs[0].add_run(value)
        set_run_font(r2, 10.2, False, TEXT)


def add_feature_table(doc):
    rows = [
        ("企业级 RAG 问答", "围绕企业知识库进行检索增强生成，支持引用片段、流式输出和多轮对话。"),
        ("知识库管理", "支持文档上传、解析、向量化索引、知识库权限控制和私有知识沉淀。"),
        ("语音助手", "提供语音入口，面向会议、客服、知识问答等场景联动 ASR、LLM 与 TTS。"),
        ("ERP 业务协同", "集成客户、合同、工单、派发、日程、审批、通知等业务模块。"),
        ("多租户与权限", "基于 JWT、RBAC、租户上下文和审计日志实现企业级安全边界。"),
        ("可运维部署", "后端基于 Spring Boot，配套 MySQL、Redis、RabbitMQ、MinIO、Qdrant 等中间件。"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    headers = ("能力模块", "说明")
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        shade_cell(cell, "123A5F")
        set_cell_borders(cell, "123A5F")
        r = cell.paragraphs[0].add_run(text)
        set_run_font(r, 10.5, True, RGBColor(255, 255, 255))
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            shade_cell(cell, "F8FAFC" if i % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(text)
            set_run_font(r, 10.1, j == 0, BLUE if j == 0 else TEXT)


def add_arch_table(doc):
    rows = [
        ("前端", "React 18、Vite、TypeScript、Tailwind CSS、shadcn/ui，提供管理台、对话页、业务工作台等界面。"),
        ("后端", "Spring Boot 3.5、Java 17、Spring Security、MyBatis-Plus，提供统一 REST API 与 SSE 流式接口。"),
        ("数据与中间件", "MySQL、Redis、RabbitMQ、MinIO、Qdrant，覆盖业务数据、缓存、异步任务、对象存储和向量检索。"),
        ("AI 能力", "可接入 DashScope 或 OpenAI 兼容模型，支持聊天模型、Embedding、语音识别和语音合成能力。"),
        ("安全机制", "JWT Token、Redis Token Store、RBAC 权限、接口限流、审计日志、敏感字段日志脱敏。"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (layer, desc) in enumerate(rows):
        c1, c2 = table.cell(i, 0), table.cell(i, 1)
        c1.width = Cm(3.3)
        c2.width = Cm(12.5)
        shade_cell(c1, "FFF7E6")
        set_cell_borders(c1)
        set_cell_borders(c2)
        set_cell_margins(c1)
        set_cell_margins(c2)
        r1 = c1.paragraphs[0].add_run(layer)
        set_run_font(r1, 10.2, True, GOLD)
        r2 = c2.paragraphs[0].add_run(desc)
        set_run_font(r2, 10.2, False, TEXT)


def build_docx():
    desktop = desktop_dir()
    output = desktop / "知语AI智能知识服务平台产品介绍.docx"
    screenshot = desktop / "知语AI-桌面端截图.png"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    for style_name in ("Normal", "Body Text"):
        if style_name in doc.styles:
            set_style_font(doc.styles[style_name], 10.5)

    paragraph(doc, COMPANY, size=12, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    paragraph(doc, TITLE, size=25, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    paragraph(doc, SUBTITLE, size=16, bold=True, color=CYAN, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    paragraph(
        doc,
        "面向企业知识管理、智能问答、语音交互与业务协同的一体化 AI 平台",
        size=11.5,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_meta_table(doc)

    heading(doc, "一、产品定位", 1)
    paragraph(
        doc,
        "知语AI智能知识服务平台以企业私有知识和业务数据为基础，通过检索增强生成、语音交互、权限控制和业务流程集成，"
        "帮助组织把分散在文档、知识库和业务系统中的信息转化为可追问、可引用、可协作的智能服务能力。",
        first_indent=True,
    )
    paragraph(
        doc,
        "平台适用于内部知识问答、政企客服、售前资料检索、项目交付支持、工单处理、合同与客户信息查询、会议日程安排等场景，"
        "既可以作为独立的企业知识助手，也可以作为现有 ERP、OA、CRM 系统的智能入口。",
        first_indent=True,
    )

    heading(doc, "二、核心能力", 1)
    add_feature_table(doc)

    heading(doc, "三、典型应用场景", 1)
    bullet(doc, "企业知识问答：员工可围绕制度、产品手册、项目资料和历史文档进行自然语言检索与追问。")
    bullet(doc, "智能客服与服务台：客服人员可通过知识库快速定位答案，工单场景可沉淀处理过程和参考资料。")
    bullet(doc, "业务工作台：客户、合同、工单、派发、日程、审批和通知统一入口，减少多系统切换成本。")
    bullet(doc, "语音交互入口：支持无登录语音助手入口，也支持登录后在对话页面进行文本和语音协同。")
    bullet(doc, "多租户管理：适合集团、园区、平台型组织按租户隔离数据、用户、角色和权限。")

    heading(doc, "四、技术架构", 1)
    paragraph(
        doc,
        "平台采用前后端分离架构。前端提供面向桌面浏览器的管理台和业务工作台；后端通过统一 API 提供认证、知识库、"
        "对话、文档处理、业务协同和审计能力；中间件层负责业务数据、缓存、异步任务、对象存储与向量检索。",
        first_indent=True,
    )
    add_arch_table(doc)

    heading(doc, "五、产品界面展示", 1)
    if screenshot.exists():
        paragraph(doc, "图 1  知语AI桌面端登录与产品入口界面", size=10, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        doc.add_picture(str(screenshot), width=Cm(15.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph(doc, "未找到桌面端截图文件，文档生成时跳过界面图片。", color=MUTED)

    heading(doc, "六、安全与运维", 1)
    bullet(doc, "认证授权：采用 JWT 登录认证、Redis Token Store、RBAC 角色权限和接口级访问控制。")
    bullet(doc, "租户隔离：通过租户上下文和数据权限机制控制跨租户访问，适配多组织部署。")
    bullet(doc, "审计与合规：关键操作通过审计日志留痕，日志输出对密码、Token、手机号等敏感字段进行脱敏。")
    bullet(doc, "接口防护：支持接口限流、统一错误响应、健康检查和 Prometheus 指标暴露。")
    bullet(doc, "部署运维：支持本地 Windows 服务、Docker Compose、反向代理和 Cloudflare Tunnel 临时公网访问。")

    heading(doc, "七、交付价值", 1)
    paragraph(
        doc,
        "知语AI不是单一聊天窗口，而是围绕企业知识、权限体系和业务流程构建的智能服务平台。"
        "它把“能回答问题”的 AI 能力落到可管理、可追溯、可部署、可持续运营的企业系统中，"
        "帮助组织提升知识复用效率、降低人工检索成本，并为后续智能客服、业务助手和语音入口扩展打下基础。",
        first_indent=True,
    )

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    paragraph(doc, "附：当前演示访问方式", size=15, bold=True, color=BLUE)
    bullet(doc, "前端本地地址：http://127.0.0.1:3000")
    bullet(doc, "后端健康检查：http://127.0.0.1:8080/api/actuator/health")
    if PUBLIC_DEMO_URL:
        bullet(doc, f"当前公网演示地址：{PUBLIC_DEMO_URL}")
        bullet(doc, "说明：临时公网隧道域名通常会在每次重新启动时变化，正式上线应使用固定子域名。")
    else:
        bullet(doc, "当前公网演示地址：未配置。可通过环境变量 ZWSK_PUBLIC_DEMO_URL 临时写入。")

    all_pages_dir = desktop / "知语AI-全部页面截图"
    page_images = sorted(all_pages_dir.glob("*.png")) if all_pages_dir.exists() else []
    if page_images:
        doc.add_section(WD_SECTION.NEW_PAGE)
        paragraph(doc, "附录：全部页面截图", size=15, bold=True, color=BLUE)
        paragraph(
            doc,
            "以下截图覆盖当前前端路由中的主要页面，用于产品介绍、验收展示和界面留档。",
            size=10.5,
            color=TEXT,
            first_indent=True,
        )
        for idx, image in enumerate(page_images, start=1):
            title = image.stem
            paragraph(doc, f"图 A-{idx}  {title}", size=10, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
            doc.add_picture(str(image), width=Cm(15.8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output)
    return output


if __name__ == "__main__":
    print(build_docx())
