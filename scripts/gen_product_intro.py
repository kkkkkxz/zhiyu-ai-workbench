# -*- coding: utf-8 -*-
"""
生成《知语 AI 智能知识服务平台产品介绍（政府与用户版）》Word 文档。

输出位置：当前用户桌面
文档定位：面向政府单位、企事业单位和公众用户的正式产品介绍材料
"""
import os
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT = "微软雅黑"
TITLE = "知语 AI 智能知识服务平台"
SUBTITLE = "政府与用户版产品介绍"
ORG = "展望数科"

BLUE = RGBColor(0x12, 0x3A, 0x5F)
CYAN = RGBColor(0x1E, 0x6D, 0x8F)
GOLD = RGBColor(0xB0, 0x84, 0x2E)
TEXT = RGBColor(0x24, 0x2B, 0x33)
MUTED = RGBColor(0x66, 0x72, 0x80)
LIGHT_BG = "F5F7FA"
PALE_BLUE = "EAF3F8"
PALE_GOLD = "FFF7E6"
TABLE_LINE = "CBD5E1"


def desktop_path() -> str:
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    if os.path.isdir(desktop):
        return desktop
    return os.path.join(os.environ.get("USERPROFILE", r"C:\Users\admin"), "Desktop")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)


def set_style_font(style, size=11):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)


def set_cell_bg(cell, color_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color_hex: str = TABLE_LINE, size: str = "6", style: str = "single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), style)
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color_hex)
        borders.append(border)
    tc_pr.append(borders)


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def para(doc, text="", size=11, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT,
         first_indent=False, before=0, after=5, line=1.35, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = line
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    size = {1: 18, 2: 14, 3: 12}.get(level, 12)
    color = {1: BLUE, 2: CYAN, 3: TEXT}.get(level, TEXT)
    before = {1: 15, 2: 10, 3: 6}.get(level, 6)
    after = {1: 8, 2: 5, 3: 4}.get(level, 4)
    p = para(doc, text, size=size, bold=True, color=color, before=before, after=after, line=1.15)
    if level == 1:
        add_bottom_border(p, "D8E2EA")
    return p


def add_bottom_border(paragraph, color_hex: str):
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color_hex)
    border.append(bottom)
    p_pr.append(border)


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run("• " + text)
    set_run_font(run, size=size, color=TEXT)
    return p


def numbered(doc, index, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.52)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(f"{index}. {text}")
    set_run_font(run, size=size, color=TEXT)
    return p


def small_label(doc, text, color=GOLD):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True, color=color)
    return p


def add_callout(doc, title, body, bg=PALE_BLUE, accent=CYAN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    set_cell_bg(cell, bg)
    set_cell_borders(cell, "D6E5EE", "8")
    set_cell_margins(cell, top=170, bottom=170, start=220, end=220)
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    r1 = p1.add_run(title)
    set_run_font(r1, size=11.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.28
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=TEXT)
    para(doc, "", after=2)


def add_meta_table(doc):
    rows = [
        ("文档名称", f"{TITLE}产品介绍"),
        ("适用对象", "政府部门、事业单位、国有企业、行业客户、公众服务用户"),
        ("文档性质", "正式产品介绍与应用场景说明"),
        ("版本日期", datetime.now().strftime("%Y年%m月%d日")),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (k, v) in enumerate(rows):
        c1, c2 = table.cell(i, 0), table.cell(i, 1)
        c1.width = Cm(3.2)
        c2.width = Cm(12.8)
        set_cell_bg(c1, "EFF6FB")
        set_cell_borders(c1)
        set_cell_borders(c2)
        set_cell_margins(c1)
        set_cell_margins(c2)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        set_run_font(r1, size=10.5, bold=True, color=BLUE)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        set_run_font(r2, size=10.5, color=TEXT)
    para(doc, "", after=4)


def add_cards(doc, cards, cols=2):
    rows = (len(cards) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    idx = 0
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.width = Cm(8)
            set_cell_bg(cell, "FFFFFF")
            set_cell_borders(cell, "D9E2EA", "8")
            set_cell_margins(cell, top=180, bottom=180, start=190, end=190)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if idx >= len(cards):
                cell.text = ""
                idx += 1
                continue
            title, body = cards[idx]
            p1 = cell.paragraphs[0]
            p1.paragraph_format.space_after = Pt(5)
            r1 = p1.add_run(title)
            set_run_font(r1, size=11.2, bold=True, color=BLUE)
            p2 = cell.add_paragraph()
            p2.paragraph_format.line_spacing = 1.22
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(body)
            set_run_font(r2, size=9.8, color=TEXT)
            idx += 1
    para(doc, "", after=3)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths or [16 / len(headers)] * len(headers)

    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = Cm(widths[i])
        set_cell_bg(cell, "123A5F")
        set_cell_borders(cell, "123A5F", "8")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for r_i, row in enumerate(rows, 1):
        for c_i, value in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.width = Cm(widths[c_i])
            set_cell_bg(cell, "FFFFFF" if r_i % 2 else "F7FAFC")
            set_cell_borders(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            text_run = p.add_run(value)
            set_run_font(text_run, size=9.6, color=TEXT)
    para(doc, "", after=3)
    return table


def add_screenshot_placeholder(doc, title, description):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    head = table.cell(0, 0)
    body = table.cell(1, 0)
    head.width = Cm(16)
    body.width = Cm(16)
    set_cell_bg(head, "123A5F")
    set_cell_borders(head, "123A5F", "8")
    set_cell_margins(head, top=100, bottom=100)
    p_head = head.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_head = p_head.add_run(f"截图占位：{title}")
    set_run_font(r_head, size=10.8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    body.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_bg(body, "F2F5F8")
    set_cell_borders(body, "9AA7B5", "10", style="dashed")
    set_cell_margins(body, top=260, bottom=260)
    for _ in range(2):
        p = body.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    p1 = body.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("请在此处粘贴或插入正式产品截图")
    set_run_font(r1, size=12, bold=True, color=BLUE)
    p2 = body.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(description)
    set_run_font(r2, size=9.5, italic=True, color=MUTED)
    for _ in range(2):
        p = body.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    para(doc, "", after=4)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        if header.paragraphs:
            p = header.paragraphs[0]
        else:
            p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"{TITLE} | {SUBTITLE}")
        set_run_font(r, size=8.5, color=MUTED)

        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{ORG}  ·  仅用于产品介绍、方案沟通和项目交流")
        set_run_font(r, size=8.5, color=MUTED)


def configure_document(doc):
    set_style_font(doc.styles["Normal"], 11)
    for section in doc.sections:
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)
    add_header_footer(doc)


def add_cover(doc):
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(ORG)
    set_run_font(run, size=14, bold=True, color=GOLD)
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(TITLE)
    set_run_font(run, size=30, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(SUBTITLE)
    set_run_font(run, size=18, bold=True, color=CYAN)

    add_callout(
        doc,
        "正式版定位",
        "面向政府部门、企事业单位和公众用户，说明平台在知识治理、智能问答、语音交互、数据安全和服务提效方面的建设价值。",
        bg=PALE_GOLD,
        accent=GOLD,
    )

    add_screenshot_placeholder(
        doc,
        "封面主视觉 / 产品 Logo",
        "建议放置产品 Logo、平台首页或正式宣传主视觉，宽度约 15-16 cm。",
    )

    para(doc, f"文档版本：V1.0    生成日期：{datetime.now():%Y年%m月%d日}",
         size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=0)
    doc.add_page_break()


def add_intro(doc):
    heading(doc, "一、产品概述", 1)
    para(
        doc,
        "知语 AI 智能知识服务平台是面向政务服务、企业服务和公共服务场景建设的智能化知识应用平台。"
        "平台通过知识库管理、检索增强生成（RAG）、流式文本问答和实时语音对话等能力，"
        "将政策文件、办事指南、业务手册、制度规范、常见问题和历史经验转化为可检索、可追溯、可复用的智能知识服务。",
        first_indent=True,
    )
    para(
        doc,
        "平台强调“可信、可管、可控、可用”。在用户体验层面，公众和工作人员可以像日常对话一样提问；"
        "在管理治理层面，政府和组织可对知识来源、权限范围、访问行为和服务效果进行统一管理；"
        "在安全合规层面，平台支持私有化部署、多租户隔离、权限控制、日志审计和敏感信息脱敏，满足严肃业务场景对数据安全的要求。",
        first_indent=True,
    )
    add_callout(
        doc,
        "一句话定位",
        "让政府和组织沉淀的政策、制度、手册与业务知识，变成人人可问、随时可用、来源可查的智能服务能力。",
    )
    add_meta_table(doc)
    add_screenshot_placeholder(
        doc,
        "产品首页 / 工作台总览",
        "建议展示登录后的平台首页、导航结构和核心入口，体现正式系统形象。",
    )


def add_audience_and_value(doc):
    heading(doc, "二、服务对象与建设价值", 1)
    add_cards(
        doc,
        [
            ("政府部门", "用于政策咨询、办事问答、热线辅助、窗口人员赋能、内部制度查询和跨部门知识沉淀，提高政务服务响应效率与一致性。"),
            ("企事业单位", "用于客服知识库、运营支撑、员工自助问答、培训辅助和业务规范查询，降低重复咨询成本。"),
            ("一线工作人员", "在办理业务、接待咨询、查找规范时获得即时答案，并可查看引用来源，减少翻找文件和经验依赖。"),
            ("公众与普通用户", "以自然语言或语音方式获取清晰答复，降低政策、流程、材料要求等信息的理解门槛。"),
        ],
        cols=2,
    )

    heading(doc, "2.1 核心建设价值", 2)
    add_table(
        doc,
        ["价值方向", "说明", "体现效果"],
        [
            ("服务提效", "将常见咨询、政策解读、流程说明交由智能助手先行响应。", "减少重复沟通，提升窗口、热线和线上服务效率。"),
            ("知识统一", "集中管理分散在文件、手册、问答库中的业务知识。", "统一口径，降低因信息不一致导致的服务偏差。"),
            ("答案可信", "基于知识库检索生成答案，并保留引用片段与来源。", "方便核验，避免只给结论、不知依据的问题。"),
            ("安全可控", "支持权限分级、租户隔离、日志审计、敏感信息脱敏。", "适配政府和组织内部数据管理要求。"),
            ("体验友好", "支持文本问答和语音交互，回答过程可流式呈现。", "降低使用门槛，适合高频咨询和现场服务。"),
        ],
        widths=[3.0, 7.0, 6.0],
    )


def add_features(doc):
    heading(doc, "三、核心功能体系", 1)
    heading(doc, "3.1 智能知识库管理", 2)
    para(
        doc,
        "平台支持按业务主题创建知识库，将政策文件、制度规范、产品资料、操作手册、常见问题等材料统一纳入管理。"
        "文档上传后，系统可进行解析、切片、向量化和索引，形成可被智能问答调用的知识底座。",
        first_indent=True,
    )
    bullet(doc, "支持知识库新建、查看、删除、权限范围设置和更新维护。")
    bullet(doc, "支持文档上传、解析状态查看、文档列表管理和无效文档清理。")
    bullet(doc, "支持异步索引处理，避免大文档处理阻塞前台操作。")
    add_screenshot_placeholder(
        doc,
        "知识库列表与知识库详情",
        "建议展示知识库列表、文档上传区域、文档解析状态和文档列表。",
    )

    heading(doc, "3.2 RAG 检索增强问答", 2)
    para(
        doc,
        "平台采用检索增强生成技术。用户提问后，系统先从指定知识库中检索相关片段，再结合大模型生成结构化答复。"
        "与普通聊天机器人相比，该方式更适合严肃业务场景，可围绕组织自有知识回答，并展示回答依据。",
        first_indent=True,
    )
    bullet(doc, "支持会话创建、历史消息查看、多轮追问和知识库挂载。")
    bullet(doc, "支持流式输出，答案生成过程实时呈现，减少等待感。")
    bullet(doc, "支持引用展开，便于查看答案对应的知识来源。")
    add_screenshot_placeholder(
        doc,
        "智能问答页面",
        "建议展示用户提问、AI 答复、引用来源、知识库选择和会话列表。",
    )

    heading(doc, "3.3 实时语音对话", 2)
    para(
        doc,
        "平台提供面向自然交互的语音对话能力，支持麦克风输入、语音识别、知识检索、模型推理和语音合成输出。"
        "用户无需输入长文本即可完成咨询，适合服务大厅、值班坐席、培训演示和移动办公等场景。",
        first_indent=True,
    )
    bullet(doc, "支持 WebSocket 实时语音链路，提供听取、思考、回答等状态反馈。")
    bullet(doc, "支持对回答进行打断和重置，适配真实对话节奏。")
    bullet(doc, "支持语音模式自动检索公开知识库或已授权知识库。")
    add_screenshot_placeholder(
        doc,
        "AI 语音对话界面",
        "建议展示语音球、通话状态、语音转写、AI 回答和引用来源。",
    )

    heading(doc, "3.4 角色、权限与审计", 2)
    para(
        doc,
        "平台内置用户、角色、权限、安全认证和审计能力，适合多部门、多岗位、多层级协同使用。"
        "管理者可按职责范围控制数据访问和功能操作，关键行为可留痕追溯。",
        first_indent=True,
    )
    bullet(doc, "支持账号登录、Token 管理、注销失效和异常访问控制。")
    bullet(doc, "支持 RBAC 权限模型，按角色控制功能与数据范围。")
    bullet(doc, "支持多租户数据隔离，避免不同组织或部门数据互相影响。")
    bullet(doc, "支持关键操作审计和敏感日志脱敏，便于后续检查与运维。")
    add_screenshot_placeholder(
        doc,
        "权限管理 / 审计日志",
        "建议展示用户角色权限、租户信息或审计记录页面。",
    )


def add_scenarios(doc):
    heading(doc, "四、典型应用场景", 1)
    add_table(
        doc,
        ["场景", "主要使用者", "应用方式", "预期效果"],
        [
            ("政务办事咨询", "公众、窗口人员", "将办事指南、政策文件、材料清单纳入知识库，提供在线问答。", "提升咨询响应速度，统一政策解释口径。"),
            ("政务热线辅助", "热线坐席、质检人员", "坐席咨询过程中实时检索标准答案和依据。", "缩短坐席培训周期，提高答复一致性。"),
            ("内部制度查询", "机关/单位工作人员", "把制度、流程、通知、手册转为可问答知识。", "减少内部反复询问，提高协同效率。"),
            ("企业客服知识库", "客服、客户、运营人员", "围绕产品手册、售后规则、常见问题自动答疑。", "降低人工客服压力，提升服务体验。"),
            ("培训与业务陪练", "新员工、基层人员", "围绕业务规范进行问答式学习和场景演练。", "提高学习效率，沉淀标准答复。"),
        ],
        widths=[3.0, 3.2, 5.8, 4.0],
    )
    add_screenshot_placeholder(
        doc,
        "典型场景展示",
        "可放置政务服务大厅、热线坐席辅助、企业知识库或培训场景截图。",
    )


def add_security(doc):
    heading(doc, "五、安全合规与可管可控", 1)
    para(
        doc,
        "对于政府和重点行业应用，智能化系统不仅要“能回答”，更要“可管理、可追溯、可控制”。"
        "知语 AI 在底层设计中纳入身份认证、权限管理、多租户隔离、审计留痕、限流防护、敏感信息脱敏和私有化部署能力，"
        "用于降低数据泄露、越权访问、服务滥用和管理失控等风险。",
        first_indent=True,
    )
    add_cards(
        doc,
        [
            ("身份可信", "基于账号认证与 Token 管理，支持登录、刷新、注销和异常失效控制。"),
            ("权限分级", "通过角色和权限模型控制用户可访问的功能、知识库和数据范围。"),
            ("数据隔离", "支持多租户隔离，避免不同组织、部门或业务线数据混用。"),
            ("审计留痕", "关键操作记录可追踪，便于运维检查、问题回溯和管理监督。"),
            ("日志脱敏", "密码、Token、密钥、手机号、身份证号等敏感字段在日志中自动脱敏。"),
            ("部署可控", "支持本地机房、政务云、专有云等私有化部署方式，具体以项目环境为准。"),
        ],
        cols=2,
    )
    add_callout(
        doc,
        "合规说明",
        "本材料所述能力基于当前系统功能和可配置能力进行说明。涉及等保、信创、商密、数据分类分级等专项合规要求时，应结合项目所在地政策、客户安全规范和实际部署方案开展专项评估。",
        bg="F8FAFC",
        accent=BLUE,
    )


def add_deployment(doc):
    heading(doc, "六、部署与运维能力", 1)
    para(
        doc,
        "平台采用前后端分离与标准化服务组件架构，便于在不同规模的组织环境中部署、扩展和维护。"
        "当前系统由 Web 前端、Java 后端服务、关系型数据库、缓存、消息队列、对象存储和向量数据库等组成，"
        "可根据客户环境选择本地化、专有云或容器化部署方式。",
        first_indent=True,
    )
    add_table(
        doc,
        ["组成", "说明", "作用"],
        [
            ("Web 前端", "面向用户提供登录、问答、知识库、语音和设置页面。", "提供统一操作入口。"),
            ("后端服务", "Spring Boot 服务承载认证、知识库、会话、RAG 和语音编排。", "处理业务逻辑与安全控制。"),
            ("数据库", "保存用户、角色、知识库、文档、会话、审计等业务数据。", "形成可管理的数据资产。"),
            ("向量数据库", "保存知识切片的向量索引。", "支撑语义检索和 RAG 问答。"),
            ("对象存储", "保存上传的原始文档或附件。", "保证资料可管理、可追溯。"),
            ("消息队列", "承载文档解析、向量化等异步任务。", "提升系统稳定性和响应速度。"),
        ],
        widths=[3.0, 8.0, 5.0],
    )
    add_screenshot_placeholder(
        doc,
        "部署架构 / 运维监控",
        "建议展示部署拓扑、服务状态、健康检查或运维监控界面。",
    )


def add_user_flow(doc):
    heading(doc, "七、用户使用流程", 1)
    small_label(doc, "面向管理员")
    numbered(doc, 1, "创建组织账号、配置角色权限和可访问知识库范围。")
    numbered(doc, 2, "创建业务知识库，上传政策文件、制度规范、产品手册或常见问题。")
    numbered(doc, 3, "检查文档解析和索引状态，确认知识内容可被检索。")
    numbered(doc, 4, "根据业务需要开放给指定部门、岗位、坐席或公众服务入口使用。")

    small_label(doc, "面向普通用户")
    numbered(doc, 1, "登录平台或进入授权服务入口。")
    numbered(doc, 2, "选择对应知识库或使用默认服务场景。")
    numbered(doc, 3, "通过文字或语音提出问题。")
    numbered(doc, 4, "查看 AI 答复、引用来源和补充说明，必要时继续追问。")
    add_screenshot_placeholder(
        doc,
        "用户完整使用流程",
        "建议放置 3-4 张流程截图：登录、选择知识库、提问、查看引用。",
    )


def add_roadmap_and_footer(doc):
    heading(doc, "八、后续建设建议", 1)
    para(
        doc,
        "结合政府和重点行业的实际落地需求，后续可围绕知识治理、统一入口、业务系统集成、质量评估和运维监管继续深化建设。",
        first_indent=True,
    )
    bullet(doc, "建立知识责任人机制：明确各知识库的维护部门、更新周期和审核流程。")
    bullet(doc, "完善服务入口：根据需要接入政务门户、内部办公系统、热线坐席系统或统一身份认证平台。")
    bullet(doc, "建立效果评估：统计高频问题、未命中问题、满意度和人工转接率，持续优化知识内容。")
    bullet(doc, "加强安全治理：围绕权限、审计、备份、等保和数据分类分级要求完善运行制度。")

    heading(doc, "九、结语", 1)
    para(
        doc,
        "知语 AI 智能知识服务平台以安全可控的知识库为基础，以智能问答和语音交互为入口，"
        "面向政府服务、组织协同和公众咨询提供高效、统一、可追溯的智能知识服务。"
        "通过持续沉淀知识、规范服务口径、提升响应效率，平台可为数字政府建设、企业数字化转型和公共服务能力提升提供切实支撑。",
        first_indent=True,
    )
    add_callout(
        doc,
        "截图替换提示",
        "文档中的“截图占位”均可直接替换为正式系统截图。建议使用同一环境、同一主题、同一分辨率截取，避免出现测试数据、个人信息、密钥、真实手机号等敏感内容。",
        bg=PALE_GOLD,
        accent=GOLD,
    )


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_intro(doc)
    add_audience_and_value(doc)
    add_features(doc)
    add_scenarios(doc)
    add_security(doc)
    add_deployment(doc)
    add_user_flow(doc)
    add_roadmap_and_footer(doc)

    out_path = os.path.join(desktop_path(), "知语AI智能知识服务平台-政府与用户版产品介绍.docx")
    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
