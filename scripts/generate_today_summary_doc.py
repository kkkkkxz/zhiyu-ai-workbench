from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont


DESKTOP = Path.home() / "Desktop"
STAMP = datetime.now().strftime("%Y-%m-%d")
DOCX_PATH = DESKTOP / f"今日工作总结_{STAMP}.docx"
PNG_PATH = DESKTOP / f"今日工作总结_{STAMP}_预览.png"


TITLE = "今日工作总结"
SUBTITLE = f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

SECTIONS = [
    (
        "一、今日完成",
        [
            "检查并修复了工单、流程审批、AI 动作定义、数据权限、日程提醒等核心模块中的明显业务问题。",
            "补齐了多租户边界校验，避免工单评论、审批任务、AI 配置等场景发生越权访问或跨租户修改。",
            "将 CRM 和合同相关接口改为白名单 DTO 入参，减少内部字段被前端覆盖的风险。",
        ],
    ),
    (
        "二、AI 方向结论",
        [
            "核查了租户级 AI API Key 管理链路，确认当前后端能兼容前端现有配置页面和运行时消费者。",
            "梳理了本地模型部署路径，明确推荐通过 OpenAI 兼容网关接入本地千问模型。",
            "给出了全部本地化与 85 人团队语音并发场景下的硬件和模型方案，并整理了预算区间。",
        ],
    ),
    (
        "三、验证结果",
        [
            "后端执行 .\\mvnw.cmd test，通过，31 个测试全部成功。",
            "前端执行 npm run typecheck，通过。",
            "本轮修复未引入新的编译错误或类型错误。",
        ],
    ),
    (
        "四、后续建议",
        [
            "继续完善 AI Key 管理，支持清空密钥、保留租户级模型参数以及更细粒度的 provider 配置。",
            "对剩余直接接收 Entity 或 Map 的业务接口逐步做 DTO 化收口。",
            "如推进全本地化部署，优先落地聊天、embedding、ASR、TTS，再考虑视觉能力。",
        ],
    ),
]


def set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(11)


def build_docx() -> None:
    document = Document()
    set_default_font(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(TITLE)
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(SUBTITLE)
    subtitle_run.font.size = Pt(10.5)
    subtitle_run.font.name = "Microsoft YaHei"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    document.add_paragraph("")

    for heading, bullets in SECTIONS:
        h = document.add_paragraph()
        h_run = h.add_run(heading)
        h_run.bold = True
        h_run.font.size = Pt(13)
        h_run.font.name = "Microsoft YaHei"
        h_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        for bullet in bullets:
            p = document.add_paragraph(style="List Bullet")
            run = p.add_run(bullet)
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(11)

    document.save(str(DOCX_PATH))


def get_font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int):
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        width = draw.textbbox((0, 0), trial, font=font)[2]
        if width <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def build_preview() -> None:
    width, height = 1600, 2000
    image = Image.new("RGB", (width, height), "#f7f4ee")
    draw = ImageDraw.Draw(image)

    title_font = get_font(64, bold=True)
    section_font = get_font(34, bold=True)
    body_font = get_font(28)
    small_font = get_font(22)

    draw.rectangle((80, 80, width - 80, height - 80), fill="#fffdf8", outline="#d8d1c2", width=3)
    draw.text((140, 140), TITLE, fill="#1f2937", font=title_font)
    draw.text((140, 230), SUBTITLE, fill="#6b7280", font=small_font)

    y = 320
    max_text_width = width - 280
    for heading, bullets in SECTIONS:
        draw.text((140, y), heading, fill="#8b5e34", font=section_font)
        y += 62
        for bullet in bullets:
            lines = wrap_text(draw, "• " + bullet, body_font, max_text_width)
            for line in lines:
                draw.text((160, y), line, fill="#374151", font=body_font)
                y += 40
            y += 16
        y += 10

    image.save(PNG_PATH)


def main() -> None:
    build_docx()
    build_preview()
    print(DOCX_PATH)
    print(PNG_PATH)


if __name__ == "__main__":
    main()
